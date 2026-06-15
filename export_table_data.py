"""
Export game_daily_metrics data for a specified date range to Excel
One sheet per game, with all columns preserved

Usage:
    # Interactive mode (prompts for dates):
    python export_october_data.py
    
    # Command line mode (provide dates as arguments):
    python export_october_data.py <start_date> <end_date>
    
    Dates should be in YYYYMMDD format (e.g., 20251001 for October 1, 2025)
    
    Examples:
        python export_october_data.py                    # Interactive prompts
        python export_october_data.py 20251001 20251031  # October 2025
        python export_october_data.py 20250901 20250930  # September 2025
"""

import mysql.connector
from mysql.connector import Error
from datetime import date
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import sys
import json
from db_config import get_db_config

# Database configuration
DB_CONFIG = get_db_config()

# Games to export
GAMES = [
    {'app_id': 2507950, 'name': 'Delta Force'},
    {'app_id': 2073620, 'name': 'Arena Breakout: Infinite'},
    {'app_id': 3478050, 'name': 'Road to Empress'},
    {'app_id': 3104410, 'name': 'Terminull Brigade'},
    {'app_id': 4148240, 'name': 'Road to Empress II'}
]

# Note: If no command line arguments are provided, the script will prompt for dates interactively


def get_all_columns():
    """Get all column names from game_daily_metrics table"""
    return [
        'stat_date',
        'game_name',
        'steam_app_id',
        'dau',
        'pcu',
        'unique_player',
        'new_players',
        'total_downloads',
        'd1_retention',
        'pcu_over_dau',
        'new_vs_returning_ratio',
        'median_playtime',
        'avg_playtime',
        'players_20h_plus',
        'lifetime_total_revenue',
        'daily_total_revenue',
        'lifetime_total_units',
        'daily_units',
        'daily_arpu',
        'top3_iap_share',
        'wishlist',
        'lifetime_wishlist_conversion_rate',
        'wishlist_additions',
        'wishlist_deletions',
        'wishlist_conversions',
        'top10_country_dau',
        'top10_country_downloads',
        'top10_region_downloads',
        'top10_country_revenue',
        'top10_region_revenue',
        'iap_breakdown_json'
    ]


def fetch_game_data(connection, app_id, start_date, end_date):
    """Fetch all data for a specific game and date range"""
    cursor = connection.cursor(dictionary=True)
    
    # Build query with all columns
    columns = ', '.join(get_all_columns())
    query = f"""
    SELECT {columns}
    FROM game_daily_metrics
    WHERE steam_app_id = %s 
    AND stat_date BETWEEN %s AND %s
    ORDER BY stat_date
    """
    
    cursor.execute(query, (app_id, start_date, end_date))
    results = cursor.fetchall()
    cursor.close()
    
    return results


def export_to_excel(all_data, output_filename):
    """Export data to Excel with one sheet per game"""
    wb = openpyxl.Workbook()
    
    # Remove default sheet
    if 'Sheet' in wb.sheetnames:
        wb.remove(wb['Sheet'])
    
    # Define styles
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    center_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    columns = get_all_columns()
    
    # Create a sheet for each game
    for game in GAMES:
        game_name = game['name']
        app_id = game['app_id']
        data = all_data.get(app_id, [])
        
        # Create sheet (Excel sheet names are limited to 31 characters and cannot contain : \ / ? * [ ])
        # Replace invalid characters with underscore
        sheet_name = game_name[:31] if len(game_name) > 31 else game_name
        invalid_chars = [':', '\\', '/', '?', '*', '[', ']']
        for char in invalid_chars:
            sheet_name = sheet_name.replace(char, '_')
        ws = wb.create_sheet(title=sheet_name)
        
        # Write headers
        for col_idx, col_name in enumerate(columns, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.value = col_name
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_alignment
            cell.border = border
        
        # Write data rows
        for row_idx, row_data in enumerate(data, start=2):
            for col_idx, col_name in enumerate(columns, start=1):
                cell = ws.cell(row=row_idx, column=col_idx)
                value = row_data.get(col_name)
                
                # Handle JSON columns - convert to string representation
                if col_name in ['top10_country_dau', 'top10_country_downloads', 
                               'top10_region_downloads', 'top10_country_revenue',
                               'top10_region_revenue', 'iap_breakdown_json']:
                    if value is not None:
                        try:
                            # If it's already a dict/list, convert to JSON string
                            if isinstance(value, (dict, list)):
                                cell.value = json.dumps(value, ensure_ascii=False)
                            else:
                                cell.value = str(value)
                        except:
                            cell.value = str(value)
                    else:
                        cell.value = None
                else:
                    cell.value = value
                
                cell.border = border
                # Center align numeric columns
                if col_name in ['dau', 'pcu', 'unique_player', 'new_players', 
                               'total_downloads', 'players_20h_plus', 
                               'lifetime_total_units', 'daily_units',
                               'wishlist', 'wishlist_additions', 'wishlist_deletions',
                               'wishlist_conversions']:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Auto-adjust column widths
        for col_idx, col_name in enumerate(columns, start=1):
            max_length = len(str(col_name))
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
                for cell in row:
                    if cell.value:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
            adjusted_width = min(max_length + 2, 50)  # Cap at 50 characters
            ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width
    
    # Save workbook
    wb.save(output_filename)
    print(f"Excel file saved: {output_filename}")


def parse_date(date_str):
    """Parse date from YYYYMMDD format"""
    try:
        if len(date_str) != 8:
            raise ValueError("Date must be 8 characters (YYYYMMDD)")
        year = int(date_str[0:4])
        month = int(date_str[4:6])
        day = int(date_str[6:8])
        return date(year, month, day)
    except (ValueError, IndexError) as e:
        raise ValueError(f"Invalid date format: {date_str}. Expected YYYYMMDD (e.g., 20251001). Error: {e}")


def main():
    """Main function"""
    # Parse dates from command line arguments or prompt user
    start_date = None
    end_date = None
    
    if len(sys.argv) == 1:
        # No arguments provided, prompt user for input
        print("="*60)
        print("Game Daily Metrics Export Tool")
        print("="*60)
        print("\nPlease enter the date range (YYYYMMDD format)")
        print("Example: 20251001 for October 1, 2025\n")
        
        while start_date is None:
            start_input = input("Enter start date (YYYYMMDD): ").strip()
            try:
                start_date = parse_date(start_input)
            except ValueError as e:
                print(f"Error: {e}")
                print("Please try again.\n")
        
        while end_date is None:
            end_input = input("Enter end date (YYYYMMDD): ").strip()
            try:
                end_date = parse_date(end_input)
                
                # Validate that start_date <= end_date
                if start_date > end_date:
                    print(f"Error: Start date ({start_date}) must be before or equal to end date ({end_date})")
                    print("Please try again.\n")
                    end_date = None
            except ValueError as e:
                print(f"Error: {e}")
                print("Please try again.\n")
        
        print()  # Empty line for readability
        
    elif len(sys.argv) == 2:
        print("Error: Both start_date and end_date are required.")
        print("Usage: python export_october_data.py <start_date> <end_date>")
        print("       Dates should be in YYYYMMDD format (e.g., 20251001 20251031)")
        print("\nOr run without arguments to be prompted for dates interactively.")
        sys.exit(1)
    elif len(sys.argv) >= 3:
        try:
            start_date = parse_date(sys.argv[1])
            end_date = parse_date(sys.argv[2])
            
            # Validate that start_date <= end_date
            if start_date > end_date:
                print(f"Error: Start date ({start_date}) must be before or equal to end date ({end_date})")
                sys.exit(1)
        except ValueError as e:
            print(f"Error: {e}")
            sys.exit(1)
    
    print(f"Exporting data from {start_date} to {end_date}")
    print(f"Games: {', '.join([g['name'] for g in GAMES])}")
    
    connection = None
    try:
        # Connect to database
        print("\nConnecting to database...")
        connection = mysql.connector.connect(**DB_CONFIG)
        print("Connected successfully!")
        
        # Fetch data for each game
        all_data = {}
        for game in GAMES:
            print(f"\nFetching data for {game['name']} (App ID: {game['app_id']})...")
            data = fetch_game_data(connection, game['app_id'], start_date, end_date)
            all_data[game['app_id']] = data
            print(f"  Found {len(data)} records")
        
        # Export to Excel
        # Generate filename based on date range
        start_str = start_date.strftime("%Y%m%d")
        end_str = end_date.strftime("%Y%m%d")
        output_filename = f"game_daily_metrics_{start_str}_{end_str}.xlsx"
        print(f"\nExporting to Excel: {output_filename}")
        export_to_excel(all_data, output_filename)
        
        # Summary
        print("\n" + "="*60)
        print("Export Summary")
        print("="*60)
        for game in GAMES:
            count = len(all_data.get(game['app_id'], []))
            print(f"{game['name']}: {count} records")
        print("="*60)
        
    except Error as e:
        print(f"Database error: {e}")
        sys.exit(1)
    except Exception as e:
        print(f"Error: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)
    finally:
        if connection and connection.is_connected():
            connection.close()
            print("\nDatabase connection closed.")


if __name__ == "__main__":
    main()

