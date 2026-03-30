"""
Extract Store Traffic Stats (Visits Over Time, Impressions Over Time) from a local HTML file
and export to Excel with columns: Date, Impression, Visits.

Usage:
    python extract_store_traffic_to_excel.py

The script will prompt for: game name, date range, and local HTML file path.
Output Excel is named: "{game name} {date range}.xlsx"
"""

import re
import os
import sys
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


def extract_series_data(html_content, var_name):
    """
    Extract the first series (Total) from dataViews or dataImpressions JavaScript array.
    Returns list of (date_str, value) tuples.
    """
    pattern = rf'var {var_name}\s*=\s*\[(.*?)\];'
    match = re.search(pattern, html_content, re.DOTALL)
    if not match:
        return None

    array_content = match.group(1)
    pair_pattern = r"\[\s*'(\d{4}-\d{2}-\d{2})',\s*(\d+)\s*\]"
    all_pairs = re.findall(pair_pattern, array_content)

    if not all_pairs:
        return None

    # First series = Total; series change when date goes backward or repeats
    result = []
    prev_date = None
    for date_str, value_str in all_pairs:
        if prev_date is not None and date_str <= prev_date:
            break
        prev_date = date_str
        result.append((date_str, int(value_str)))

    return result


def merge_visits_and_impressions(visits_data, impressions_data):
    """
    Merge visits and impressions by date. Returns list of (date, impressions, visits).
    Uses union of dates; missing values become None.
    """
    by_date = {}
    for date_str, value in (visits_data or []):
        by_date[date_str] = {'impressions': None, 'visits': value}
    for date_str, value in (impressions_data or []):
        if date_str not in by_date:
            by_date[date_str] = {'impressions': value, 'visits': None}
        else:
            by_date[date_str]['impressions'] = value

    sorted_dates = sorted(by_date.keys())
    return [
        (d, by_date[d]['impressions'], by_date[d]['visits'])
        for d in sorted_dates
    ]


def export_to_excel(rows, output_path):
    """Export data to Excel with Date, Impression, Visits columns."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Store Traffic Stats"

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    headers = ['Date', 'Impression', 'Visits']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col)
        cell.value = h
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row_idx, (date_str, impressions, visits) in enumerate(rows, start=2):
        ws.cell(row=row_idx, column=1, value=date_str).border = border
        ws.cell(row=row_idx, column=2, value=impressions).border = border
        ws.cell(row=row_idx, column=3, value=visits).border = border

    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 14

    wb.save(output_path)


def main():
    print("=" * 60)
    print("Store Traffic Stats - Extract to Excel")
    print("=" * 60)

    game_name = input("\nEnter the game name: ").strip()
    if not game_name:
        print("Error: No game name provided.")
        sys.exit(1)

    date_range = input("Enter the date range: ").strip()
    if not date_range:
        print("Error: No date range provided.")
        sys.exit(1)

    file_path = input("Enter the local HTML file path: ").strip()
    if not file_path:
        print("Error: No file path provided.")
        sys.exit(1)

    file_path = file_path.strip('"').strip("'")
    if not os.path.isfile(file_path):
        print(f"Error: File not found: {file_path}")
        sys.exit(1)

    print(f"\nReading: {file_path}")
    with open(file_path, 'r', encoding='utf-8', errors='replace') as f:
        html_content = f.read()

    visits_data = extract_series_data(html_content, 'dataViews')
    impressions_data = extract_series_data(html_content, 'dataImpressions')

    if not visits_data and not impressions_data:
        print("Error: Could not extract dataViews or dataImpressions from the HTML file.")
        sys.exit(1)

    rows = merge_visits_and_impressions(visits_data, impressions_data)
    if not rows:
        print("Error: No data extracted.")
        sys.exit(1)

    safe_name = re.sub(r'[\\/:*?"<>|]', '_', f"{game_name} {date_range}")
    output_filename = f"{safe_name}.xlsx"
    output_path = os.path.join(os.path.dirname(os.path.abspath(file_path)), output_filename)
    export_to_excel(rows, output_path)

    print(f"\nExported {len(rows)} rows to: {output_path}")
    print(f"Date range: {rows[0][0]} to {rows[-1][0]}")


if __name__ == "__main__":
    main()
